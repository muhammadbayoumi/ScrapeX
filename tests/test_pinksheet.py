"""worldbank-pink-sheet: the World Bank's monthly commodity prices (#1912).

The workbook fixture is the real file of 2026-10-02 with the sheets the connector does not
read removed, so the census below is the real census. The page is a window of the live one.
"""
from __future__ import annotations

import io
import sqlite3
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from scrapex import db as dbmod
from scrapex.config import MANIFEST_FILE, ExtractSpec, SourceEntry, load_manifest
from scrapex.connectors.base import ScrapedTable
from scrapex.connectors.factory import _BUILDERS, supports_history
from scrapex.connectors.pinksheet import (
    _SERIES,
    HIDDEN_SOURCE,
    PinkSheetConnector,
    change_points,
    find_workbook_url,
    series_name,
)
from scrapex.ingest import ingest_payloads
from scrapex.rowspec import COMMODITY_PRICE, ENRICHMENT, RowBuilder, RowView
from scrapex.tax import resolve
from scrapex.vocab import ConnectorFamily, ExtractKind, ExtractScope, VatMode

FX = Path(__file__).parent / "fixtures"
WORKBOOK = (FX / "pinksheet_monthly_2026-10-02.xlsx").read_bytes()
PAGE = (FX / "pinksheet_page.html").read_text(encoding="utf-8")
LINK = ("https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/"
        "related/CMO-Historical-Data-Monthly.xlsx")
_SHORT: bytes | None = None


class _Resp:
    def __init__(self, text="", content=b""):
        self.text, self.content = text, content


class _Stub:
    """The two requests a check costs, and nothing else."""

    def __init__(self, page=PAGE, workbook=WORKBOOK):
        self.page, self.workbook, self.urls = page, workbook, []

    def get(self, url, **kwargs):
        self.urls.append(url)
        if url.endswith("/en/research/commodity-markets"):
            return _Resp(text=self.page)
        assert url == LINK, f"the connector went somewhere else: {url}"
        return _Resp(content=self.workbook)


def make_entry() -> SourceEntry:
    return SourceEntry.model_validate(dict(
        source_key="WORLDBANK_PINK_SHEET", source_name="World Bank Pink Sheet",
        base_url="https://www.worldbank.org", family="worldbank-pink-sheet",
        cadence="weekly", authority="official", currency="USD", vat_mode="excl",
        extract=[ExtractSpec(kind=ExtractKind.COMMODITY_PRICE, scope=ExtractScope.TARGETED),
                 ExtractSpec(kind=ExtractKind.ENRICHMENT, scope=ExtractScope.TARGETED)],
    ))


def tables(*, history=False, stub=None):
    connector = PinkSheetConnector(stub or _Stub(), history=history)
    return list(connector.fetch(make_entry()))


def prices(table) -> list[dict]:
    view = RowView(COMMODITY_PRICE, table.header)
    return [view.as_dict(r) for r in table.rows]


def details(table) -> list[dict]:
    view = RowView(ENRICHMENT, table.header)
    return [view.as_dict(r) for r in table.rows]


def _shorter() -> openpyxl.Workbook:
    """The real workbook less 1960M04 to 2023M12: every shape, a thirtieth of the wait."""
    book = openpyxl.load_workbook(io.BytesIO(WORKBOOK))
    sheet = book["Monthly Prices"]
    first, last = row_of(sheet, "1960M04"), row_of(sheet, "2023M12")
    sheet.delete_rows(first, last - first + 1)
    return book


def edited(change) -> bytes:
    """The real workbook with one thing changed, as a later edition might change it."""
    global _SHORT
    if _SHORT is None:
        out = io.BytesIO()
        _shorter().save(out)
        _SHORT = out.getvalue()
    book = openpyxl.load_workbook(io.BytesIO(_SHORT))
    change(book)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def column_of(sheet, name: str) -> int:
    for cell in sheet[5]:
        if cell.value is not None and series_name(str(cell.value)) == name:
            return cell.column
    raise AssertionError(name)


def row_of(sheet, period: str) -> int:
    for cell in sheet["A"]:
        if cell.value == period:
            return cell.row
    raise AssertionError(period)


@pytest.fixture(scope="module")
def weekly():
    return tables()


@pytest.fixture(scope="module")
def full():
    return tables(history=True)


# ---- the landing page --------------------------------------------------------

def test_the_workbook_address_is_read_from_the_page():
    assert find_workbook_url(PAGE) == LINK


def test_a_page_that_stopped_linking_it_is_loud():
    with pytest.raises(ValueError, match="links 0"):
        find_workbook_url("<a href='x'>Monthly prices</a>")


def test_two_different_workbooks_on_the_page_are_loud_not_a_guess():
    other = LINK.replace("74e8be41", "aaaaaaaa")
    with pytest.raises(ValueError, match="links 2"):
        find_workbook_url(f'<a href="{LINK}"></a><a href="{other}"></a>')


def test_the_same_link_twice_is_one_workbook():
    assert find_workbook_url(f'<a href="{LINK}"></a><a href="{LINK}"></a>') == LINK


def test_a_link_to_another_host_is_never_followed():
    with pytest.raises(ValueError, match="links 0"):
        find_workbook_url('<a href="https://evil.example/en/CMO-Historical-Data-Monthly.xlsx"></a>')


def test_a_check_costs_two_requests():
    stub = _Stub()
    list(PinkSheetConnector(stub).fetch(make_entry()))
    assert len(stub.urls) == 2


# ---- the census, against the study's own counts --------------------------------

def test_every_priced_series_comes_through_and_the_index_does_not(full):
    rows = prices(full[0])
    assert {r["material_key"] for r in rows} == {v[0] for v in _SERIES.values()}
    assert len(_SERIES) == 70
    assert "NATURAL_GAS_INDEX" not in {r["material_key"] for r in rows}
    assert not any(r["unit"].startswith("2010") for r in rows)
    assert not any("Natural gas index" in w for w in full[0].warnings),         "the index is left out on purpose, so it is not a new series to review"


def test_history_is_the_change_points_of_all_seventy_plus_the_hidden_values(full):
    rows = prices(full[0])
    published = [r for r in rows if r["provenance"] == "reported"
                 and r["official_source_name"] != HIDDEN_SOURCE]
    hidden = [r for r in rows if r["official_source_name"] == HIDDEN_SOURCE]
    # #1912 counted 40,787 with Rice Thai 25%'s run of zeros as one point. It is refused here.
    assert len(published) == 40_786
    assert len(hidden) == 138


def test_a_normal_run_reads_a_year_and_not_the_whole_history(weekly, full):
    reported = [r for r in prices(weekly[0]) if r["provenance"] == "reported"]
    assert reported and min(r["as_of_date"] for r in reported) == "2025-10-01"
    assert len(reported) < 1_000 < len(prices(full[0]))


def test_the_newest_month_is_the_observed_price(weekly):
    rows = {r["material_key"]: r for r in prices(weekly[0]) if r["provenance"] == "observed"}
    assert rows["COPPER"]["price"] == "14474"
    assert rows["COPPER"]["source_date"] == "2026-09-01"
    assert rows["COPPER"]["as_of_date"] == ""


def test_a_late_series_is_still_current_and_a_dead_one_is_not(weekly):
    rows = {r["material_key"]: r for r in prices(weekly[0]) if r["provenance"] == "observed"}
    assert rows["TOBACCO_US"]["source_date"] == "2026-07-01"
    for dead in ("BARLEY", "SORGHUM", "SHRIMP"):
        assert dead not in rows, f"{dead} last published long ago and is not a current price"
    assert len(rows) == 67
    said = " ".join(weekly[0].warnings)
    assert "Barley: no current price, last published 2020M08" in said
    assert "Shrimps, Mexican: no current price, last published 2023M10" in said


# ---- what a number is ------------------------------------------------------------

def test_a_row_says_what_the_study_found(weekly):
    row = next(r for r in prices(weekly[0]) if r["material_key"] == "COPPER")
    assert row["currency"] == "USD" and row["unit"] == "tonne"
    assert row["country_code_alpha2"] == "*"
    assert row["consumer_segment"] == "wholesale"
    assert row["tax_included"] == "0"
    assert row["price_basis"] == "original" and row["original_price"] == ""
    assert row["official_source_link"] == LINK
    assert "Bloomberg" in row["official_source_name"]


def test_plywood_is_stated_in_dollars_with_the_cents_beside_it(weekly):
    row = next(r for r in prices(weekly[0])
               if r["material_key"] == "PLYWOOD" and r["provenance"] == "observed")
    assert row["price"] == "3.491" and row["original_price"] == "349.1"
    assert row["currency"] == "USD" and row["unit"] == "sheet"
    assert row["price_basis"] == "converted"


def test_iron_ore_is_per_dry_ton_and_the_workbook_word_survives(weekly):
    row = next(r for r in prices(weekly[0]) if r["material_key"] == "IRON_ORE")
    assert row["unit"] == "dmt"
    word = next(d for d in details(weekly[1])
                if d["external_product_id"] == "IRON_ORE" and d["attribute_code"] == "unit_as_published")
    assert word["raw_value"] == "($/dmtu)"


@pytest.mark.parametrize("key, unit", [
    ("CRUDE_OIL_BRENT", "bbl"), ("NATURAL_GAS_US", "mmbtu"), ("LOGS_CAMEROON", "m3"),
    ("GOLD", "troy oz"), ("COCOA", "kg"), ("ALUMINUM", "tonne"),
])
def test_every_kind_of_unit(weekly, key, unit):
    assert next(r for r in prices(weekly[0]) if r["material_key"] == key)["unit"] == unit


def test_a_flag_on_the_header_is_not_part_of_the_series():
    assert series_name("Coal, South African **") == "Coal, South African"
    assert series_name("Copper") == "Copper"


# ---- holes in the sheet ------------------------------------------------------------

def test_a_zero_and_an_excel_error_are_holes_and_are_said_out_loud(full):
    rice = {r["as_of_date"]: r["price"] for r in prices(full[0])
            if r["material_key"] == "RICE_THAI_25" and r["provenance"] == "reported"}
    for month in ("2008-03-01", "2008-04-01", "2008-05-01", "2008-06-01", "2008-02-01", "1988-08-01"):
        assert month not in rice
    assert all(Decimal(p) > 0 for p in rice.values())
    assert any("Rice, Thai 25%: 4 cell(s) of 0 and 4 Excel error(s)" in w for w in full[0].warnings)


def test_an_absent_cell_is_silent_because_it_is_ordinary(full):
    assert not any("DAP" in w for w in full[0].warnings), "DAP's '...' before 1967 is not news"
    dap = [r["as_of_date"] for r in prices(full[0])
           if r["material_key"] == "DAP" and r["provenance"] == "reported"]
    assert min(dap) == "1967-01-01"


def test_change_points_keep_the_first_and_each_move():
    values = {"a": Decimal(1), "b": Decimal(1), "c": Decimal(2), "e": Decimal(2), "f": Decimal(1)}
    assert list(change_points(values, ["a", "b", "c", "d", "e", "f"])) == [
        ("a", Decimal(1)), ("c", Decimal(2)), ("f", Decimal(1))]


def test_change_points_of_nothing_is_nothing():
    assert list(change_points({}, ["a"])) == []


# ---- the hidden sheet: his word, «تتبع اسعار طالما مذكورة» --------------------------

def test_the_hidden_grains_are_taken_and_marked(full):
    rows = [r for r in prices(full[0]) if r["official_source_name"] == HIDDEN_SOURCE]
    assert {r["material_key"] for r in rows} == {"BARLEY", "SORGHUM"}
    assert all(r["provenance"] == "reported" for r in rows)
    newest = max((r for r in rows if r["material_key"] == "BARLEY"), key=lambda r: r["as_of_date"])
    assert newest["as_of_date"] == "2026-07-01" and newest["price"] == "162.8"
    public = [r for r in prices(full[0]) if r["material_key"] == "BARLEY"
              and r["official_source_name"] != HIDDEN_SOURCE]
    assert max(r["as_of_date"] for r in public) < "2020-09-01"


def test_a_hidden_value_is_not_the_older_edition_of_a_published_one(full):
    """The sheet's `Value mismatch` rows hold what July was before it was revised."""
    lng = {r["as_of_date"]: r for r in prices(full[0]) if r["material_key"] == "LNG_JAPAN"}
    assert lng["2026-07-01"]["price"] == "13.85"
    assert lng["2026-07-01"]["official_source_name"] != HIDDEN_SOURCE


def test_the_older_edition_rows_are_not_even_considered(full):
    assert not any("hidden value(s) ignored" in w for w in full[0].warnings), \
        "only `Blank vs number` rows are candidates, and on the real file none has a public cell"


def test_a_published_cell_beats_the_hidden_value_for_the_same_month():
    def publish_barley(book):
        sheet = book["Monthly Prices"]
        sheet.cell(row_of(sheet, "2026M07"), column_of(sheet, "Barley"), 999.5)
    table = tables(history=True, stub=_Stub(workbook=edited(publish_barley)))[0]
    july = [r for r in prices(table) if r["material_key"] == "BARLEY"
            and r["as_of_date"] == "2026-07-01"]
    assert [r["price"] for r in july] == ["999.5"]
    assert july[0]["official_source_name"] != HIDDEN_SOURCE
    assert any("hidden value(s) ignored" in w for w in table.warnings)


def test_when_the_hidden_sheet_goes_the_run_says_so_and_carries_on():
    def remove(book):
        del book["Mismatch Details"]
    table = tables(history=True, stub=_Stub(workbook=edited(remove)))[0]
    assert any("hidden sheet" in w and "gone" in w for w in table.warnings)
    assert not table.defects, "a QA sheet going away does not degrade the published prices"
    assert not any(r["official_source_name"] == HIDDEN_SOURCE for r in prices(table))
    assert "COPPER" in {r["material_key"] for r in prices(table)}


def test_a_hidden_sheet_with_other_columns_is_read_as_nothing_and_said():
    def rename(book):
        book["Mismatch Details"]["F4"] = "Something else"
    table = tables(history=True, stub=_Stub(workbook=edited(rename)))[0]
    assert any("changed its columns" in w for w in table.warnings)
    assert not any(r["official_source_name"] == HIDDEN_SOURCE for r in prices(table))


# ---- every parse asserts its shape ---------------------------------------------------

@pytest.mark.parametrize("sheet", ["Monthly Prices", "Description"])
def test_a_renamed_sheet_is_loud(sheet):
    def rename(book):
        book[sheet].title = "Something else"
    with pytest.raises(ValueError, match=sheet):
        tables(stub=_Stub(workbook=edited(rename)))


def test_a_sheet_that_stopped_saying_nominal_dollars_is_loud():
    def rewrite(book):
        book["Monthly Prices"]["A2"] = "monthly prices in real 2010 US dollars"
    with pytest.raises(ValueError, match="nominal US dollars"):
        tables(stub=_Stub(workbook=edited(rewrite)))


def test_a_first_column_with_no_months_is_loud():
    def rewrite(book):
        for cell in book["Monthly Prices"]["A"]:
            if cell.row > 6:
                cell.value = "2026-09"
    with pytest.raises(ValueError, match="YYYYMmm"):
        tables(stub=_Stub(workbook=edited(rewrite)))


def test_a_word_where_a_month_belongs_is_loud():
    def rewrite(book):
        sheet = book["Monthly Prices"]
        sheet.cell(row_of(sheet, "2025M01"), 1, "Notes")
    with pytest.raises(ValueError, match="where a month belongs"):
        tables(stub=_Stub(workbook=edited(rewrite)))


def test_months_out_of_order_are_loud():
    def rewrite(book):
        sheet = book["Monthly Prices"]
        sheet.cell(row_of(sheet, "2025M02"), 1, "1999M01")
    with pytest.raises(ValueError, match="out of order"):
        tables(stub=_Stub(workbook=edited(rewrite)))


def test_a_word_in_a_price_cell_is_loud_not_a_missing_month():
    def rewrite(book):
        sheet = book["Monthly Prices"]
        sheet.cell(row_of(sheet, "2025M01"), column_of(sheet, "Copper"), "discontinued")
    with pytest.raises(ValueError, match="no numeric content"):
        tables(stub=_Stub(workbook=edited(rewrite)))


def test_an_unknown_unit_is_loud():
    def rewrite(book):
        sheet = book["Monthly Prices"]
        sheet.cell(6, column_of(sheet, "Copper"), "(EUR/mt)")
    with pytest.raises(ValueError, match="price unit"):
        tables(stub=_Stub(workbook=edited(rewrite)))


def test_a_new_series_is_reported_and_not_guessed():
    def add(book):
        sheet = book["Monthly Prices"]
        sheet.cell(5, sheet.max_column + 1, "Cement, world")
    table = tables(stub=_Stub(workbook=edited(add)))[0]
    assert any("Cement, world" in w for w in table.warnings)
    assert not table.defects


def test_a_renamed_series_is_a_defect_on_the_run():
    def rename(book):
        sheet = book["Monthly Prices"]
        sheet.cell(5, column_of(sheet, "Copper"), "Copper, LME")
    table = tables(stub=_Stub(workbook=edited(rename)))[0]
    assert any("Copper" in d for d in table.defects)
    assert any("Copper, LME" in w for w in table.warnings)
    assert "COPPER" not in {r["material_key"] for r in prices(table)}


def test_a_series_twice_is_loud():
    def double(book):
        sheet = book["Monthly Prices"]
        sheet.cell(5, column_of(sheet, "Silver"), "Copper")
    with pytest.raises(ValueError, match="twice"):
        tables(stub=_Stub(workbook=edited(double)))


def test_a_workbook_that_lost_most_of_its_series_is_loud():
    def thin(book):
        sheet = book["Monthly Prices"]
        for cell in list(sheet[5])[1:]:
            if cell.column > 20:
                cell.value = None
    with pytest.raises(ValueError, match="known series"):
        tables(stub=_Stub(workbook=edited(thin)))


def test_a_body_far_past_the_usual_size_is_refused_before_it_is_opened():
    with pytest.raises(ValueError, match="bytes"):
        tables(stub=_Stub(workbook=b"x" * 25_000_001))


def test_a_body_that_is_not_a_workbook_is_loud():
    with pytest.raises(Exception, match="(?i)zip"):
        tables(stub=_Stub(workbook=b"<html>a login page</html>"))


def test_a_description_row_gone_is_said_and_the_prices_stay():
    def remove(book):
        sheet = book["Description"]
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("Copper (LME)"):
                    cell.value = None
    table, facts = tables(stub=_Stub(workbook=edited(remove)))
    assert any("no Description row for 'Copper'" in w for w in table.warnings)
    assert "COPPER" in {r["material_key"] for r in prices(table)}
    assert "market_terms" not in {d["attribute_code"] for d in details_of(facts, "COPPER")}


def details_of(table, key: str) -> list[dict]:
    return [d for d in details(table) if d["external_product_id"] == key]


# ---- the attributes his answers ask for ------------------------------------------------

def test_every_series_says_it_is_a_monthly_average_and_where_it_trades(full):
    keys = {v[0] for v in _SERIES.values()}
    for code in ("price_period", "market_terms", "commodity_group", "unit_as_published"):
        have = {d["external_product_id"] for d in details(full[1]) if d["attribute_code"] == code}
        assert have == keys, f"{code} missing on {sorted(keys - have)}"
    copper = {d["attribute_code"]: d["raw_value"] for d in details_of(full[1], "COPPER")}
    assert copper["price_period"].startswith("Monthly average")
    assert copper["market_terms"].startswith("Copper (LME)")
    assert copper["commodity_group"] == "Metals and minerals"


def test_each_description_row_went_to_its_own_series(full):
    """Rows are found by how they begin, so two series must not share one."""
    terms = [d["raw_value"] for d in details(full[1]) if d["attribute_code"] == "market_terms"]
    assert len(terms) == len(set(terms)) == 70


# ---- into the warehouse ------------------------------------------------------------------

@pytest.fixture
def conn():
    connection: sqlite3.Connection = dbmod.connect(":memory:")
    dbmod.migrate(connection)
    yield connection
    connection.close()


def ingest(conn, found):
    return ingest_payloads(conn, make_entry(), [t.to_payload() for t in found])


def test_a_weekly_run_lands_with_no_errors(conn, weekly):
    result = ingest(conn, weekly)
    assert not result.errors, result.errors[:3]
    assert result.rejected_out_of_scope == 0, "a detail with no product of its own was refused"
    offers = conn.execute("SELECT COUNT(*) FROM source_offer").fetchone()[0]
    # Shrimp stopped in 2023 and nothing in the last year names it, hidden sheet included.
    assert offers == 69


def test_a_series_with_nothing_this_run_sends_no_details_to_be_refused(weekly):
    assert details_of(weekly[1], "SHRIMP") == []
    assert "SHRIMP" not in {r["material_key"] for r in prices(weekly[0])}


def test_every_offer_is_wholesale_and_worldwide(conn, weekly):
    ingest(conn, weekly)
    assert [tuple(r) for r in conn.execute(
        "SELECT DISTINCT customer_segment, country_code_alpha2, currency, tax_included "
        "FROM source_offer")] == [("wholesale", "*", "USD", 0)]


def test_a_full_history_lands_with_no_errors(conn, full):
    result = ingest(conn, full)
    assert not result.errors, result.errors[:3]
    assert result.rejected_out_of_scope == 0
    assert conn.execute("SELECT COUNT(*) FROM source_offer").fetchone()[0] == 70
    reported = conn.execute(
        "SELECT COUNT(*) FROM price_observation WHERE provenance = 'reported'").fetchone()[0]
    assert reported == 40_786 + 138


def test_a_second_run_of_the_same_edition_adds_nothing_reported(conn, weekly):
    ingest(conn, weekly)
    before = conn.execute(
        "SELECT COUNT(*) FROM price_observation WHERE provenance = 'reported'").fetchone()[0]
    ingest(conn, weekly)
    assert conn.execute(
        "SELECT COUNT(*) FROM price_observation WHERE provenance = 'reported'").fetchone()[0] == before


def test_a_revised_month_lands_beside_the_old_value_and_overwrites_nothing(conn, weekly):
    ingest(conn, weekly)

    def revise(book):
        sheet = book["Monthly Prices"]
        sheet.cell(row_of(sheet, "2026M07"), column_of(sheet, "Copper"), 13000)
    ingest(conn, tables(stub=_Stub(workbook=edited(revise))))
    july = conn.execute(
        "SELECT po.price FROM price_observation po "
        "JOIN source_offer so ON so.offer_id = po.offer_id "
        "JOIN source_variant sv ON sv.source_variant_id = so.source_variant_id "
        "JOIN source_product sp ON sp.source_product_id = sv.source_product_id "
        "WHERE sp.external_product_id = 'COPPER' AND po.provenance = 'reported' "
        "AND po.business_date = '2026-07-01' ORDER BY po.price").fetchall()
    assert [r[0] for r in july] == [13000.0, 13543.0]


def test_the_units_have_names_in_both_languages(conn, weekly):
    ingest(conn, weekly)
    rows = conn.execute("SELECT unit_code, name, name_ar FROM selling_unit").fetchall()
    assert {r[0] for r in rows} == {"tonne", "kg", "bbl", "mmbtu", "m3", "troy oz", "dmt", "sheet"}
    assert all(r[1] and r[2] for r in rows), rows


def test_a_retail_row_still_lands_retail(conn):
    """The widening reaches only a row that states a segment."""
    entry = SourceEntry.model_validate(dict(
        source_key="ARAMCO_FUEL_SA", source_name="Saudi Aramco", base_url="https://www.aramco.com",
        family="aramco-fuel-page", cadence="monthly", authority="official", currency="SAR",
        default_region="SA", vat_mode="incl",
        extract=[ExtractSpec(kind=ExtractKind.COMMODITY_PRICE, scope=ExtractScope.LATEST_ONLY,
                             materials=["DIESEL"], regions=["SA"])],
    ))
    builder = RowBuilder(COMMODITY_PRICE)
    row = builder.row(material_key="DIESEL", country_code_alpha2="SA", currency="SAR", unit="liter",
                      tax_included="1", price="1.79", provenance="observed")
    table = ScrapedTable(entry.source_key, COMMODITY_PRICE.kind, "u", builder.header, [row])
    ingest_payloads(conn, entry, [table.to_payload()])
    assert [r[0] for r in conn.execute("SELECT customer_segment FROM source_offer")] == ["retail"]


def test_a_segment_the_warehouse_does_not_know_is_refused_out_loud(conn):
    entry = make_entry()
    builder = RowBuilder(COMMODITY_PRICE)
    row = builder.row(material_key="COPPER", country_code_alpha2="*", currency="USD", unit="kg",
                      tax_included="0", price="1", provenance="observed", consumer_segment="household")
    table = ScrapedTable(entry.source_key, COMMODITY_PRICE.kind, "u", builder.header, [row])
    result = ingest_payloads(conn, entry, [table.to_payload()])
    assert result.errors and "CHECK" in result.errors[0]
    assert conn.execute("SELECT COUNT(*) FROM source_offer").fetchone()[0] == 0


# ---- the manifest and the registry -------------------------------------------------------

def test_the_manifest_entry_ships_off_and_obeying():
    entry = load_manifest(MANIFEST_FILE).get("WORLDBANK_PINK_SHEET")
    assert entry.family == ConnectorFamily.WORLDBANK_PINK_SHEET
    assert entry.active is False, "no source is made active without his word"
    assert str(entry.robots) == "obey"
    assert entry.crawl_pace_s == 2.5
    assert entry.cadence.value == "weekly"
    assert entry.vat_mode == VatMode.EXCLUSIVE
    assert entry.tax == [], "the workbook says nothing about tax, so no rule may claim it"


def test_the_family_is_registered_and_offers_a_history_backfill():
    assert ConnectorFamily.WORLDBANK_PINK_SHEET in _BUILDERS
    assert supports_history(ConnectorFamily.WORLDBANK_PINK_SHEET)


def test_tax_reads_unverified_because_nothing_is_claimed():
    state = resolve({}, "*", "COPPER").for_row(False)
    assert state.evidence == "unknown"
    assert state.short_label() == "Unverified"
